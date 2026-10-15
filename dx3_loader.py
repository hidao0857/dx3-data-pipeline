"""ダブルクロス3rd のキャラクターシート（テキスト）を読み込み、
キャラ管理DB（xlsx）の各シートに1行ずつ書き込むスクリプト。

処理の流れ（ETL）
  1. Extract   : テキストファイルを読み込む
  2. Transform : 「■エフェクト■」などのセクションごとに分解し、各シートの列に合わせた行データに変換する
  3. Load      : xlsx の空いている行に値を書き込む（自動計算の列には触らない）
  4. Master    : エフェクト・アイテム・覚醒・衝動のマスタに、まだ無い名称を追記する
  5. Check     : テキストに書かれていた HP・行動値などと、xlsx の計算式で出る値が一致するか確かめる

使い方
  python dx3_loader.py
    → 作業フォルダの inputTxt フォルダにある .txt をすべて読み込み、DX3_キャラ管理DB.xlsx に書き込む
  python dx3_loader.py -o 出力.xlsx                  （元の xlsx は変えずに別名で保存）
  python dx3_loader.py --input-dir 別フォルダ --xlsx 別の.xlsx
"""

import argparse
import datetime
import re
import sys
from copy import copy
from pathlib import Path

import openpyxl

BASE_DIR = Path(__file__).resolve().parent  # 作業フォルダ（このスクリプトがあるフォルダ）

# 覚醒・衝動マスタを新しく作るときに最初から入れておく侵蝕値（ルールブック参照）。
# 以降は xlsx の「覚醒マスタ」「衝動マスタ」シートの値を使う。シート側に直接追記してもよい。
AWAKENING_EROSION = {"感染": 14}
IMPULSE_EROSION = {"破壊": 16}

# マスタシート：シート名 → 列の並び（1列目がキー。キーは名称のみで、キャラが違っても同じ名称は1行にまとめる）
MASTER_SHEETS = {
    "エフェクトマスタ": ["エフェクト名", "タイミング", "技能", "難易度", "対象", "射程", "侵蝕値", "制限",
                     "シンドローム", "効果", "出典", "初出キャラID", "登録日"],
    "アイテムマスタ": ["名称", "種別", "技能", "命中", "攻撃力", "ガード値", "射程", "ドッジ", "行動", "装甲値",
                    "常備化", "経験点", "効果", "初出キャラID", "登録日"],
    "覚醒マスタ": ["覚醒", "侵蝕値", "メモ", "初出キャラID", "登録日"],
    "衝動マスタ": ["衝動", "侵蝕値", "メモ", "初出キャラID", "登録日"],
}
MASTER_SEEDS = {"覚醒マスタ": AWAKENING_EROSION, "衝動マスタ": IMPULSE_EROSION}

# テキストの略記 → xlsx のプルダウン（リストシート）の表記
TIMING_MAP = {
    "オート": "オートアクション",
    "メジャー": "メジャーアクション",
    "マイナー": "マイナーアクション",
    "リアクション": "リアクション",
    "セットアップ": "セットアッププロセス",
    "イニシアチブ": "イニシアチブプロセス",
    "クリンナップ": "クリンナッププロセス",
    "常時": "常時",
}
BREED_BY_COUNT = {1: "ピュア", 2: "クロス", 3: "トライ"}
COMMON_EFFECTS = {"リザレクト", "ワーディング"}
NONE_MARKS = {"", "―", "-", "－", "—"}


# ---------------------------------------------------------------------------
# 1. Extract / 2. Transform
# ---------------------------------------------------------------------------

def to_int(text):
    """'37' や '8ｍ' や '30％' から数値だけを取り出す。数字が無ければ None。"""
    if text is None:
        return None
    m = re.search(r"-?\d+", text.translate(str.maketrans("０１２３４５６７８９", "0123456789")))
    return int(m.group()) if m else None


def blank_to_none(text):
    """空欄や「―」は None に、数字だけの文字列は数値に変換する。"""
    text = (text or "").strip()
    if text in NONE_MARKS:
        return None
    return int(text) if re.fullmatch(r"-?\d+", text) else text


def split_sections(lines):
    """'■エフェクト■' のような見出しで行を分ける。見出しより前の行は 'ヘッダー' に入れる。"""
    sections = {"ヘッダー": []}
    current = "ヘッダー"
    for line in lines:
        m = re.fullmatch(r"\s*■(.+?)■\s*", line)
        if m:
            current = m.group(1)
            sections[current] = []
        else:
            sections[current].append(line)
    return sections


def parse_key_values(lines):
    """'キー：値' 形式の行を辞書にする（全角スペースでの桁揃えは取り除く）。"""
    result = {}
    for line in lines:
        if "：" in line:
            key, value = line.split("：", 1)
            result[key.replace("　", "").strip()] = value.strip()
    return result


def parse_slash_table(lines):
    """'名前 / Lv / ...' 形式の表を、1行目を見出しとした辞書のリストにする。"""
    rows = [[c.strip() for c in line.split("/")] for line in lines if "/" in line]
    if not rows:
        return []
    header = rows[0]
    return [dict(zip(header, r)) for r in rows[1:]]


def parse_charasheet(text):
    lines = text.replace("\r\n", "\n").split("\n")
    sec = split_sections(lines)
    head = parse_key_values(sec["ヘッダー"])
    life = parse_key_values(sec.get("ライフパス", []))
    others = sec.get("その他", [])

    syndromes = [s for s in re.split(r"[／/×・、,\s]+", head.get("シンドローム", "")) if s]
    syndromes = [("ブラム＝ストーカー" if s in ("ブラム=ストーカー", "ブラムストーカー") else s) for s in syndromes]

    # 能力値と技能：【肉体】：8 (内成長：1) / 〈白兵〉：4 / 判定 8r+4
    abilities, skills, stats = [], [], {}
    current_ability = None
    for line in sec.get("能力値と技能", []):
        m = re.match(r"【(.+?)】：\s*(\d+)\s*\(内成長：(\d+)\)", line)
        if m:
            current_ability = m.group(1)
            abilities.append({"能力値": current_ability, "合計": int(m.group(2)), "成長": int(m.group(3))})
            continue
        m = re.match(r"〈(.+?)〉：\s*(\d+)", line)
        if m and current_ability:
            skills.append({"能力値": current_ability, "技能": m.group(1), "Lv": int(m.group(2))})
            continue
        m = re.match(r"【(.+?)】\s*(.+)", line)
        if m:
            stats[m.group(1).replace("　", "").strip()] = to_int(m.group(2))

    # エフェクト
    effects = []
    for row in parse_slash_table(sec.get("エフェクト", [])):
        name = row["《エフェクト名》"].strip("《》").rstrip("：:")
        text_syndrome = re.search(r"【(.+?)】", row.get("効果", ""))  # 効果欄の「シンドローム【オルクス】」
        text_syndrome = text_syndrome.group(1) if text_syndrome and text_syndrome.group(1) in syndromes else None
        if name in COMMON_EFFECTS:
            syndrome = "共通"
        elif len(syndromes) == 1:  # ピュアブリードなら全エフェクトがそのシンドローム
            syndrome = syndromes[0]
        else:
            syndrome = text_syndrome
        if name == "コンセントレイト" and syndrome:  # 《コンセントレイト：》 → コンセントレイト:キュマイラ
            name = f"コンセントレイト:{syndrome}"
        effects.append({
            "エフェクト名": name,
            "SL": to_int(row.get("Lv")),
            "タイミング": TIMING_MAP.get(row.get("タイミング"), blank_to_none(row.get("タイミング"))),
            "難易度": blank_to_none(row.get("難易度")),
            "対象": blank_to_none(row.get("対象")),
            "射程": blank_to_none(row.get("射程")),
            "侵蝕値": blank_to_none(row.get("侵蝕値")),
            "制限": blank_to_none(row.get("制限")),
            "シンドローム": syndrome,
            "効果": blank_to_none(row.get("効果")),
        })

    # アイテム（・武器 / ・防具 / ・その他 の3つの小見出し）
    items, kind, buffer = [], None, []
    item_kinds = {"武器", "防具", "一般", "ヴィークル", "コネ", "その他"}  # xlsx のプルダウンの選択肢

    def flush():
        for row in parse_slash_table(buffer):
            item_kind, name, note = kind, blank_to_none(row.get("名前")), blank_to_none(row.get("解説"))
            if kind == "その他":
                # 「・その他」の中は、行ごとの種別（コネ・使い捨てなど）を見て振り分ける
                row_kind = blank_to_none(row.get("種別"))
                if row_kind in item_kinds:
                    item_kind = row_kind
                elif row_kind:
                    item_kind = "一般"
                    note = f"[{row_kind}] {note or ''}".strip()
                if item_kind == "コネ" and isinstance(name, str):
                    name = re.sub(r"^コネ[：:]\s*", "", name)  # 「コネ：UGN幹部」→「UGN幹部」
            skill = blank_to_none(row.get("技能"))
            items.append({
                "種別": item_kind if item_kind in item_kinds else "その他",
                "名称": name,
                "技能": skill.strip("〈〉") if isinstance(skill, str) else skill,
                "命中": to_int(row.get("命中")),
                "攻撃力": to_int(row.get("攻撃力")),
                "ガード値": to_int(row.get("ガード値")),
                "射程": blank_to_none(row.get("射程")),
                "ドッジ": to_int(row.get("ドッジ")),
                "行動": to_int(row.get("行動値")),
                "装甲値": to_int(row.get("装甲値")),
                "常備化": to_int(row.get("常備化")),
                "経験点": to_int(row.get("経験点")),
                "効果": note,
            })

    for line in sec.get("アイテム", []):
        m = re.match(r"・(.+)", line.strip())
        if m:
            flush()
            kind, buffer = m.group(1).strip(), []
        elif line.strip().startswith("【"):
            key = re.match(r"【(.+?)】", line.strip()).group(1)
            stats[key] = to_int(line)
        else:
            buffer.append(line)
    flush()

    combos = parse_combos(parse_slash_table(sec.get("コンボ", [])))
    roises = parse_slash_table(sec.get("ロイス", []))
    memories = parse_slash_table(sec.get("メモリー", []))

    # その他セクションは自由記述なので、全体をそのままメモに入れる
    memo = "\n".join(others).strip() or None

    return {
        "キャラクター": {
            "キャラ名": head.get("キャラクター名"),
            "コードネーム": head.get("コードネーム"),
            "年齢": to_int(head.get("年齢")),
            "性別": head.get("性別"),
            "身長": to_int(head.get("身長")),
            "体重": to_int(head.get("体重")),
            "ワークス": head.get("ワークス"),
            "カヴァー": head.get("カヴァー"),
            "ブリード": BREED_BY_COUNT.get(len(syndromes)),
            "シンドローム": syndromes,
            # 出自・経験・邂逅はシートの書き方がキャラごとに違うため取り込まず、空欄で登録する（xlsx で手入力）
            "出自": None,
            "経験": None,
            "邂逅/欲望": None,
            "覚醒": life.get("覚醒"),
            "衝動": life.get("衝動"),
            "侵蝕基本値": stats.get("侵蝕基本値"),
            "HP": stats.get("ＨＰ") or stats.get("HP"),
            "行動値": stats.get("行動値"),
            "戦闘移動": stats.get("戦闘移動"),
            "常備化ポイント": stats.get("常備化ポイント"),
            "財産ポイント": stats.get("財産ポイント"),
            "メモ": memo,
        },
        "能力値": abilities,
        "技能": skills,
        "エフェクト": effects,
        "アイテム": items,
        "コンボ": combos,
        "ロイス": roises,
        "メモリー": memories,
    }


def sum_terms(text):
    """'7+4' や '4+5' のような足し算の式を計算する。"""
    return sum(int(n) for n in re.findall(r"[+-]?\d+", text or ""))


def parse_hit(text):
    """命中の式 '(7+4)dx+(4+4)@8' を、ダイス数・達成値修正・C値に分ける。"""
    m = re.fullmatch(r"\(?([\d+\-]+)\)?dx(?:\+\(?([\d+\-]+)\)?)?(?:@(\d+))?", (text or "").replace(" ", ""))
    if not m:
        return None, None, None
    return sum_terms(m.group(1)), sum_terms(m.group(2)), int(m.group(3)) if m.group(3) else None


def parse_combos(rows):
    """同じコンボ名の「100%未満」「100%以上」の2行を、1つのコンボ（xlsx の1行）にまとめる。"""
    combos = {}
    for row in rows:
        name = blank_to_none(row.get("コンボ名"))
        if not name:
            continue
        combo = combos.setdefault(name, {
            "コンボ名": name,
            "構成エフェクト": blank_to_none(row.get("組み合わせ")),
            "技能": blank_to_none(row.get("技能")),
            "対象": blank_to_none(row.get("対象")),
            "射程": blank_to_none(row.get("射程")),
            "侵蝕値": blank_to_none(row.get("侵蝕値")),
            "メモ": blank_to_none(row.get("効果")),
        })
        dice, bonus, critical = parse_hit(row.get("命中"))
        values = {"ダイス": dice, "C値": critical, "攻撃力": to_int(row.get("攻撃力")), "達成値修正": bonus}
        condition = row.get("条件", "")
        # 条件が無い行は、100%未満・以上の両方に同じ値を入れる
        brackets = [b for b in ("100%未満", "100%以上") if b in condition] or ["100%未満", "100%以上"]
        for bracket in brackets:
            for key, value in values.items():
                combo[f"{bracket} {key}"] = value
    return list(combos.values())


def split_erosion(awakening, impulse, base, awakening_master, impulse_master):
    """覚醒・衝動それぞれの侵蝕値を決める。マスタに無い側は侵蝕基本値から逆算する。
    戻り値：(覚醒の侵蝕値, 衝動の侵蝕値, 逆算したのはどちらか)"""
    a = (awakening_master.get(awakening) or {}).get("侵蝕値")
    i = (impulse_master.get(impulse) or {}).get("侵蝕値")
    derived = None
    if a is None and i is not None and base is not None:
        a, derived = base - i, "覚醒"
    if i is None and a is not None and base is not None:
        i, derived = base - a, "衝動"
    if a is None and i is None:
        print(f"  [注意] 覚醒「{awakening}」と衝動「{impulse}」の侵蝕値がどちらもマスタに無く、分けられません。"
              f"xlsx の覚醒マスタか衝動マスタに侵蝕値を入れてから再実行してください（侵蝕基本値 {base}）")
    return a, i, derived


# ---------------------------------------------------------------------------
# 3. Load / 4. Master
# ---------------------------------------------------------------------------

def header_index(ws):
    """見出し名 → 列番号 の辞書。列の位置が変わっても見出し名で書き込めるようにする。"""
    return {c.value: c.column for c in ws[1] if c.value}


def first_empty_row(ws):
    for row in range(2, ws.max_row + 1):
        if ws.cell(row, 1).value in (None, ""):
            return row
    return ws.max_row + 1


def write_row(ws, values):
    """キャラIDが空いている最初の行に、見出し名で指定した列だけ書き込む（計算式の列はそのまま残す）。"""
    cols = header_index(ws)
    row = first_empty_row(ws)
    for key, value in values.items():
        if value is not None:
            ws.cell(row, cols[key]).value = value
    return row


def find_or_new_id(wb, name):
    """同じキャラ名がすでに登録されていればそのIDを返し、古い行を消す（再実行しても重複しない）。"""
    ws = wb["キャラクター"]
    cols = header_index(ws)
    ids = []
    for row in range(2, ws.max_row + 1):
        cid = ws.cell(row, 1).value
        if cid:
            ids.append(cid)
            if ws.cell(row, cols["キャラ名"]).value == name:
                clear_character(wb, cid)
                return cid
    numbers = [int(i[1:]) for i in ids if re.fullmatch(r"C\d+", str(i))]
    return f"C{max(numbers, default=0) + 1:03d}"


def ensure_master_sheets(wb):
    """マスタシートが無ければ作る。見出しの書式は「エフェクト」シートの見出しに合わせる。"""
    template = wb["エフェクト"]
    for title, columns in MASTER_SHEETS.items():
        if title in wb.sheetnames:
            continue
        ws = wb.create_sheet(title)
        for col, name in enumerate(columns, start=1):
            cell = ws.cell(1, col, name)
            src = template.cell(1, 1)
            cell.font, cell.fill, cell.border, cell.alignment = (copy(src.font), copy(src.fill),
                                                                 copy(src.border), copy(src.alignment))
            letter = cell.column_letter
            ws.column_dimensions[letter].width = 40 if name in ("効果", "メモ") else 14
        ws.freeze_panes = "A2"
        for name, erosion in MASTER_SEEDS.get(title, {}).items():
            write_row(ws, {columns[0]: name, "侵蝕値": erosion, "メモ": "ルールブック"})


def read_master(wb, title):
    """マスタシートを {名称: {列名: 値}} の辞書で読む。"""
    ws = wb[title]
    header = [c.value for c in ws[1]]
    return {row[0]: dict(zip(header, row)) for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}


def register_masters(wb, data, cid, erosion):
    """各マスタに、まだ登録されていない名称だけを追記する（すでにある名称は上書きしない）。"""
    today = datetime.date.today().isoformat()
    ch = data["キャラクター"]
    awakening_e, impulse_e, derived = erosion
    candidates = {
        "エフェクトマスタ": [{k: e[k] for k in MASTER_SHEETS["エフェクトマスタ"] if k in e} for e in data["エフェクト"]],
        "アイテムマスタ": [{k: i[k] for k in MASTER_SHEETS["アイテムマスタ"] if k in i} for i in data["アイテム"]],
        "覚醒マスタ": [{"覚醒": ch["覚醒"], "侵蝕値": awakening_e,
                    "メモ": "侵蝕基本値から逆算" if derived == "覚醒" else None}],
        "衝動マスタ": [{"衝動": ch["衝動"], "侵蝕値": impulse_e,
                    "メモ": "侵蝕基本値から逆算" if derived == "衝動" else None}],
    }
    added = {}
    for title, records in candidates.items():
        key = MASTER_SHEETS[title][0]
        known = set(read_master(wb, title))
        added[title] = []
        for record in records:
            name = record.get(key)
            if name and name not in known:
                write_row(wb[title], {**record, "初出キャラID": cid, "登録日": today})
                known.add(name)
                added[title].append(name)
    return added


def clear_character(wb, cid):
    """指定IDの行から、入力セル（計算式でないセル）だけを空にする。"""
    for ws in wb.worksheets:
        if ws.title in ("使い方", "リスト") or ws.title in MASTER_SHEETS:
            continue
        for row in ws.iter_rows(min_row=2):
            if row[0].value == cid:
                for cell in row:
                    if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                        cell.value = None


def load_to_workbook(data, wb):
    ch = data["キャラクター"]
    cid = find_or_new_id(wb, ch["キャラ名"])
    erosion = split_erosion(ch["覚醒"], ch["衝動"], ch["侵蝕基本値"],
                            read_master(wb, "覚醒マスタ"), read_master(wb, "衝動マスタ"))
    awakening_e, impulse_e, _ = erosion
    syn = ch["シンドローム"] + [None, None, None]

    write_row(wb["キャラクター"], {
        "キャラID": cid, "キャラ名": ch["キャラ名"], "コードネーム": ch["コードネーム"],
        "年齢": ch["年齢"], "性別": ch["性別"], "身長": ch["身長"], "体重": ch["体重"],
        "ワークス": ch["ワークス"], "カヴァー": ch["カヴァー"], "ブリード": ch["ブリード"],
        "シンドローム1": syn[0], "シンドローム2": syn[1], "シンドローム3": syn[2],
        "出自": ch["出自"], "経験": ch["経験"], "邂逅/欲望": ch["邂逅/欲望"],
        "覚醒": ch["覚醒"], "覚醒 侵蝕値": awakening_e, "衝動": ch["衝動"], "衝動 侵蝕値": impulse_e,
        "現在侵蝕率": ch["侵蝕基本値"], "現在HP": ch["HP"],
        "財産ポイント": ch["財産ポイント"], "メモ": ch["メモ"],
    })

    for a in data["能力値"]:
        # テキストには「合計」と「内成長」しか無いので、残りはまとめて「その他」に入れる
        write_row(wb["能力値"], {
            "キャラID": cid, "能力値": a["能力値"], "成長": a["成長"],
            "その他": a["合計"] - a["成長"], "メモ": "シンドローム/ワークスの内訳はシートに記載なし",
        })

    for s in data["技能"]:
        if s["Lv"] > 0:  # Lv0 の技能は「持っていない」のと同じなので登録しない
            write_row(wb["技能"], {"キャラID": cid, **s})

    for e in data["エフェクト"]:
        write_row(wb["エフェクト"], {"キャラID": cid, **e})

    for item in data["アイテム"]:
        write_row(wb["アイテム"], {"キャラID": cid, **item})

    for c in data["コンボ"]:
        write_row(wb["コンボ"], {"キャラID": cid, **c})

    for r in data["ロイス"]:
        relation = blank_to_none(r.get("関係"))
        write_row(wb["ロイス"], {
            "キャラID": cid, "対象": blank_to_none(r.get("名前")), "関係": relation,
            "感情P": blank_to_none(r.get("ポジティブ")), "感情N": blank_to_none(r.get("ネガティブ")),
            "状態": blank_to_none(r.get("状態")) or "ロイス",
            "Dロイス": "はい" if relation == "Dロイス" else "いいえ",
            "メモ": blank_to_none(r.get("メモ")),
        })

    for m in data["メモリー"]:
        write_row(wb["メモリー"], {
            "キャラID": cid, "対象": m.get("名前"), "関係": m.get("関係"),
            "感情": blank_to_none(m.get("感情")), "エピソード": blank_to_none(m.get("メモ")),
        })

    return cid, erosion


# ---------------------------------------------------------------------------
# 5. Check
# ---------------------------------------------------------------------------

def expected_values(data):
    """テキストに書かれていた値を、xlsx の計算式と同じルールで検算する。"""
    ab = {a["能力値"]: a["合計"] for a in data["能力値"]}
    procure = sum(s["Lv"] for s in data["技能"] if s["技能"] == "調達")
    ch = data["キャラクター"]
    checks = [
        ("HP最大値", ch["HP"], ab["肉体"] * 2 + ab["精神"] + 20),
        ("行動値", ch["行動値"], ab["感覚"] * 2 + ab["精神"]),
        ("戦闘移動", ch["戦闘移動"], ab["感覚"] * 2 + ab["精神"] + 5),
        ("常備化ポイント", ch["常備化ポイント"], ab["社会"] * 2 + procure * 2),
    ]
    ok = True
    for name, in_text, calculated in checks:
        mark = "OK" if in_text == calculated else "不一致"
        ok &= in_text == calculated
        print(f"  {name:<8} テキスト={in_text}  計算={calculated}  {mark}")
    return ok


def read_sheet_text(path):
    """UTF-8 で読めなければ、Windows のメモ帳などで保存された Shift_JIS（cp932）として読み直す。"""
    try:
        return path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        return path.read_text(encoding="cp932")


def validate(data):
    """キャラシートとして読めたかを確かめる。足りなければ、xlsx に書き込む前に止める。"""
    ch = data["キャラクター"]
    problems = []
    if not ch["キャラ名"]:
        problems.append("「キャラクター名：」の行がない")
    abilities = sorted(a["能力値"] for a in data["能力値"])
    if abilities != sorted(["肉体", "感覚", "精神", "社会"]):
        problems.append(f"能力値が肉体・感覚・精神・社会の4つそろっていない（読めたもの: {abilities}）")
    if problems:
        raise ValueError("キャラシートの形式ではないようです：" + "、".join(problems))


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input-dir", type=Path, default=BASE_DIR / "inputTxt", help="キャラシートの .txt を置くフォルダ")
    parser.add_argument("--xlsx", type=Path, default=BASE_DIR / "DX3_キャラ管理DB.xlsx", help="書き込み先のキャラ管理DB")
    parser.add_argument("-o", "--output", type=Path, help="出力先 xlsx（省略時は --xlsx を上書き）")
    args = parser.parse_args()

    if not args.input_dir.is_dir():
        sys.exit(f"[エラー] フォルダが見つかりません: {args.input_dir}")
    texts = sorted(args.input_dir.glob("*.txt"))  # ファイル名順に取り込む
    if not texts:
        sys.exit(f"[エラー] {args.input_dir} に .txt がありません")
    print(f"[対象] {args.input_dir} の {len(texts)}件: {'、'.join(p.name for p in texts)}")

    wb = openpyxl.load_workbook(args.xlsx)
    ensure_master_sheets(wb)
    all_ok, failed = True, []
    for path in texts:
        print(f"[読込] {path.name}")
        try:
            data = parse_charasheet(read_sheet_text(path))
            validate(data)
            cid, erosion = load_to_workbook(data, wb)
        except Exception as e:  # 1件読めなくても、残りのファイルは取り込む
            print(f"  [失敗] {type(e).__name__}: {e}")
            failed.append(path.name)
            continue
        print(f"  → {cid} {data['キャラクター']['キャラ名']}：能力値{len(data['能力値'])}行 / "
              f"技能{sum(1 for s in data['技能'] if s['Lv'] > 0)}行 / エフェクト{len(data['エフェクト'])}行 / "
              f"コンボ{len(data['コンボ'])}行 / アイテム{len(data['アイテム'])}行 / ロイス{len(data['ロイス'])}行")
        added = register_masters(wb, data, cid, erosion)
        print("[マスタ] 新しく登録した名称")
        for title, names in added.items():
            print(f"  {title:<8} {len(names)}件 {'、'.join(names)}")
        print("[検算] テキストの値と、xlsx の計算式で出る値の比較")
        all_ok &= expected_values(data)

    out = args.output or args.xlsx
    wb.save(out)
    print(f"[保存] {out}")
    print(f"[結果] 成功 {len(texts) - len(failed)}件 / 失敗 {len(failed)}件" + (f"（{'、'.join(failed)}）" if failed else ""))
    sys.exit(0 if all_ok and not failed else 1)


if __name__ == "__main__":
    main()
